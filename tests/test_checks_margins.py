"""
Тесты проверок полей страницы (margins).

Запуск:
  pytest tests/test_checks_margins.py -v
  pytest tests/test_checks_margins.py --cov=vkr_checker --cov-report=term-missing
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm

from vkr_checker.parser import load_document, DocumentModel


FIXTURES = Path(__file__).parent / "fixtures"


def save_doc(doc: Document, tmp_path: Path, name: str = "test.docx") -> Path:
    """Сохраняет документ и возвращает путь к нему."""
    p = tmp_path / name
    doc.save(str(p))
    return p


class TestMarginsBasics:
    """Базовые тесты полей страницы."""

    def test_correct_margins(self, tmp_path):
        """Проверка корректных полей (левое 3см, правое 1см, верхнее 2см, нижнее 2см)."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        
        doc.add_paragraph("Текст с корректными полями")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert abs(section.left_margin.cm - 3.0) < 0.05
        assert abs(section.right_margin.cm - 1.0) < 0.05
        assert abs(section.top_margin.cm - 2.0) < 0.05
        assert abs(section.bottom_margin.cm - 2.0) < 0.05

    def test_wrong_left_margin(self, tmp_path):
        """Проверка неправильного левого поля (должно быть 3см)."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(2.0)  # Неправильно
        section.right_margin = Cm(1.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        
        doc.add_paragraph("Текст с неправильным левым полем")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert abs(section.left_margin.cm - 2.0) < 0.05
        # Левое поле не равно 3см
        assert section.left_margin.cm != 3.0

    def test_wrong_right_margin(self, tmp_path):
        """Проверка неправильного правого поля (должно быть 1см)."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)  # Неправильно
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        
        doc.add_paragraph("Текст с неправильным правым полем")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert abs(section.right_margin.cm - 2.0) < 0.05
        # Правое поле не равно 1см
        assert section.right_margin.cm != 1.0

    def test_all_margins_wrong(self, tmp_path):
        """Проверка когда все поля неправильные."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(2.0)   # должно быть 3.0
        section.right_margin = Cm(2.0)  # должно быть 1.0
        section.top_margin = Cm(3.0)    # должно быть 2.0
        section.bottom_margin = Cm(3.0) # должно быть 2.0
        
        doc.add_paragraph("Текст с неправильными полями")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert section.left_margin.cm != 3.0
        assert section.right_margin.cm != 1.0
        assert section.top_margin.cm != 2.0
        assert section.bottom_margin.cm != 2.0


class TestMarginsMultipleSections:
    """Тесты полей для документов с несколькими секциями."""

    def test_multiple_sections_same_margins(self, tmp_path):
        """Документ с несколькими секциями с одинаковыми полями."""
        doc = Document()
        
        # Первая секция
        section1 = doc.sections[0]
        section1.left_margin = Cm(3.0)
        section1.right_margin = Cm(1.0)
        section1.top_margin = Cm(2.0)
        section1.bottom_margin = Cm(2.0)
        
        doc.add_paragraph("Секция 1")
        
        # Добавляем разрыв раздела
        doc.add_section()
        
        # Вторая секция с теми же полями
        section2 = doc.sections[1]
        section2.left_margin = Cm(3.0)
        section2.right_margin = Cm(1.0)
        section2.top_margin = Cm(2.0)
        section2.bottom_margin = Cm(2.0)
        
        doc.add_paragraph("Секция 2")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        assert len(model.doc.sections) == 2
        
        for section in model.doc.sections:
            assert abs(section.left_margin.cm - 3.0) < 0.05
            assert abs(section.right_margin.cm - 1.0) < 0.05
            assert abs(section.top_margin.cm - 2.0) < 0.05
            assert abs(section.bottom_margin.cm - 2.0) < 0.05

    def test_multiple_sections_different_margins(self, tmp_path):
        """Документ с несколькими секциями с разными полями."""
        doc = Document()
        
        # Первая секция с корректными полями
        section1 = doc.sections[0]
        section1.left_margin = Cm(3.0)
        section1.right_margin = Cm(1.0)
        section1.top_margin = Cm(2.0)
        section1.bottom_margin = Cm(2.0)
        
        doc.add_paragraph("Секция 1")
        
        # Добавляем разрыв раздела
        doc.add_section()
        
        # Вторая секция с другими полями
        section2 = doc.sections[1]
        section2.left_margin = Cm(2.5)
        section2.right_margin = Cm(1.5)
        section2.top_margin = Cm(2.5)
        section2.bottom_margin = Cm(2.5)
        
        doc.add_paragraph("Секция 2")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        assert len(model.doc.sections) == 2
        
        # Проверяем первую секцию
        s1 = model.doc.sections[0]
        assert abs(s1.left_margin.cm - 3.0) < 0.05
        
        # Проверяем вторую секцию
        s2 = model.doc.sections[1]
        assert abs(s2.left_margin.cm - 2.5) < 0.05


class TestMarginsEdgeCases:
    """Граничные случаи проверки полей."""

    def test_zero_margins(self, tmp_path):
        """Проверка полей равных нулю."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(0)
        section.right_margin = Cm(0)
        section.top_margin = Cm(0)
        section.bottom_margin = Cm(0)
        
        doc.add_paragraph("Текст с нулевыми полями")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert section.left_margin.cm == 0
        assert section.right_margin.cm == 0
        assert section.top_margin.cm == 0
        assert section.bottom_margin.cm == 0

    def test_very_large_margins(self, tmp_path):
        """Проверка очень больших полей."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(10.0)
        section.right_margin = Cm(10.0)
        section.top_margin = Cm(10.0)
        section.bottom_margin = Cm(10.0)
        
        doc.add_paragraph("Текст с большими полями")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert abs(section.left_margin.cm - 10.0) < 0.05
        assert abs(section.right_margin.cm - 10.0) < 0.05
        assert abs(section.top_margin.cm - 10.0) < 0.05
        assert abs(section.bottom_margin.cm - 10.0) < 0.05

    def test_gutter_margin(self, tmp_path):
        """Проверка поля для переплёта (gutter)."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.gutter = Cm(0.5)  # Поле для переплёта
        
        doc.add_paragraph("Текст с полем для переплёта")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert abs(section.gutter.cm - 0.5) < 0.05


class TestMarginsTolerance:
    """Тесты на точность проверки полей."""

    @pytest.mark.parametrize("left,right,top,bottom", [
        (3.0, 1.0, 2.0, 2.0),
        (3.05, 1.0, 2.0, 2.0),  # Небольшое отклонение левого
        (3.0, 1.05, 2.0, 2.0),  # Небольшое отклонение правого
        (3.0, 1.0, 2.05, 2.0),  # Небольшое отклонение верхнего
        (3.0, 1.0, 2.0, 2.05),  # Небольшое отклонение нижнего
    ])
    def test_margins_precision(self, tmp_path, left, right, top, bottom):
        """Проверка точности сохранения значений полей."""
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        
        doc.add_paragraph(f"Тест точности: {left}/{right}/{top}/{bottom}")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        section = model.doc.sections[0]
        assert abs(section.left_margin.cm - left) < 0.05
        assert abs(section.right_margin.cm - right) < 0.05
        assert abs(section.top_margin.cm - top) < 0.05
        assert abs(section.bottom_margin.cm - bottom) < 0.05
