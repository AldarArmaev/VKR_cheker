"""
Интеграционные тесты проверок с использованием конфигурации rules.yaml.

Эти тесты:
1. Загружают реальную конфигурацию из config/rules.yaml
2. Создают документы с нарушениями
3. Запускают реальные проверки (FontCheck, MarginsCheck и др.)
4. Проверяют, что возвращаются Issue с правильными rule_id, severity и сообщениями

Запуск:
  pytest tests/test_checks_integration.py -v
"""
from __future__ import annotations

from pathlib import Path
import yaml
import pytest
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from vkr_checker.parser import load_document
from vkr_checker.style_resolver import StyleResolver
from vkr_checker.checks.base import CheckResult, Severity, Issue
from vkr_checker.checks.fonts import FontsCheck
from vkr_checker.checks.margins import MarginsCheck
from vkr_checker.checks.spacing import SpacingCheck
from vkr_checker.checks.indents import IndentsCheck


FIXTURES = Path(__file__).parent / "fixtures"
CONFIG_PATH = Path(__file__).parent.parent / "config" / "rules.yaml"


def load_config() -> dict:
    """Загружает конфигурацию из rules.yaml."""
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def save_doc(doc: Document, tmp_path: Path, name: str = "test.docx") -> Path:
    """Сохраняет документ и возвращает путь к нему."""
    p = tmp_path / name
    doc.save(str(p))
    return p


class TestFontsCheckIntegration:
    """Интеграционные тесты проверки шрифтов с выдачей Issue."""

    def test_font_name_violation_creates_issue(self, tmp_path):
        """Проверка что неправильный шрифт создаёт Issue."""
        config = load_config()
        required_font = config["fonts"]["required_font"]  # Times New Roman
        
        doc = Document()
        # Добавляем Введение чтобы проверки работали
        doc.add_paragraph("Введение")
        # Добавляем несколько параграфов с неправильным шрифтом для прохождения дедупликации
        for i in range(5):
            para = doc.add_paragraph(f"Текст шрифтом Arial {i}")
            para.runs[0].font.name = "Arial"  # Нарушение
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Проверяем что есть ошибки
        assert result.error_count > 0
        
        # Проверяем что есть Issue с правильным rule_id
        font_issues = [i for i in result.issues if i.rule_id == "font_name"]
        assert len(font_issues) >= 1
        
        issue = font_issues[0]
        assert issue.severity == Severity.ERROR
        assert "Arial" in issue.message
        assert required_font in issue.message

    def test_correct_font_no_issues(self, tmp_path):
        """Проверка что правильный шрифт не создаёт Issue."""
        config = load_config()
        
        doc = Document()
        para = doc.add_paragraph("Текст шрифтом Times New Roman")
        para.runs[0].font.name = "Times New Roman"
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Не должно быть ошибок шрифта
        font_issues = [i for i in result.issues if i.rule_id == "font_name"]
        assert len(font_issues) == 0

    def test_font_size_violation_creates_issue(self, tmp_path):
        """Проверка что неправильный размер шрифта создаёт Issue."""
        config = load_config()
        required_size = config["fonts"]["sizes"]["body"]  # 14
        
        doc = Document()
        # Добавляем Введение чтобы проверки работали
        doc.add_paragraph("Введение")
        # Добавляем обычный текст после Введения
        para = doc.add_paragraph("Текст кеглем 12pt")
        para.runs[0].font.size = Pt(12)  # Нарушение: должно быть 14
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Проверяем что есть ошибки размера
        size_issues = [i for i in result.issues if i.rule_id == "font_size"]
        assert len(size_issues) >= 1
        
        issue = size_issues[0]
        assert issue.severity == Severity.ERROR
        # Проверяем что сообщение содержит информацию о нарушении размера
        assert "пт" in issue.message

    def test_font_size_14pt_no_issues(self, tmp_path):
        """Проверка что размер 14pt не создаёт Issue."""
        config = load_config()
        
        doc = Document()
        para = doc.add_paragraph("Текст кеглем 14pt")
        para.runs[0].font.size = Pt(14)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        size_issues = [i for i in result.issues if i.rule_id == "font_size"]
        assert len(size_issues) == 0

    def test_bold_in_body_creates_warning(self, tmp_path):
        """Проверка что жирный текст в теле создаёт Warning."""
        config = load_config()
        
        doc = Document()
        # Добавляем Введение чтобы войти в основной текст
        doc.add_paragraph("Введение")
        # Добавляем Главу 1 чтобы выйти из Введения (где bold разрешён)
        doc.add_paragraph("Глава 1")
        # Обычный параграф с жирным текстом (нарушение) - после Глава 1
        para = doc.add_paragraph("Жирный текст в теле")
        para.runs[0].font.bold = True
        para.runs[0].font.name = "Times New Roman"
        para.runs[0].font.size = Pt(14)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Проверяем предупреждения о жирном тексте
        bold_issues = [i for i in result.issues if i.rule_id == "bold_in_body"]
        assert len(bold_issues) >= 1
        
        issue = bold_issues[0]
        assert issue.severity == Severity.WARNING

    def test_bold_in_heading_allowed(self, tmp_path):
        """Проверка что жирный текст в заголовках разрешён."""
        config = load_config()
        
        doc = Document()
        doc.add_paragraph("Введение")
        # Заголовок главы - жирный разрешён
        heading = doc.add_paragraph("Глава 1")
        heading.runs[0].font.bold = True
        heading.runs[0].font.name = "Times New Roman"
        heading.runs[0].font.size = Pt(14)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Не должно быть предупреждений о жирном в заголовках
        bold_issues = [i for i in result.issues if i.rule_id == "bold_in_body"]
        # Жирный в заголовке "Глава 1" должен быть разрешён
        assert len(bold_issues) == 0


class TestMarginsCheckIntegration:
    """Интеграционные тесты проверки полей с выдачей Issue."""

    def test_correct_margins_no_issues(self, tmp_path):
        """Проверка что корректные поля не создают Issue."""
        config = load_config()
        expected = config["margins_cm"]
        
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(expected["left"])
        section.right_margin = Cm(expected["right"])
        section.top_margin = Cm(expected["top"])
        section.bottom_margin = Cm(expected["bottom"])
        
        doc.add_paragraph("Текст с корректными полями")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = MarginsCheck(rules=config)
        result = check.run(model, resolver)
        
        assert result.error_count == 0
        assert len(result.issues) == 0

    def test_wrong_left_margin_creates_issue(self, tmp_path):
        """Проверка что неправильное левое поле создаёт Issue."""
        config = load_config()
        expected = config["margins_cm"]
        tol = expected.get("tolerance", 0.15)
        
        doc = Document()
        section = doc.sections[0]
        # Неправильное левое поле (выходит за tolerance)
        section.left_margin = Cm(expected["left"] + tol + 0.5)
        section.right_margin = Cm(expected["right"])
        section.top_margin = Cm(expected["top"])
        section.bottom_margin = Cm(expected["bottom"])
        
        doc.add_paragraph("Текст с неправильным левым полем")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = MarginsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Должна быть ошибка
        assert result.error_count > 0
        
        margin_issues = [i for i in result.issues if "margin_left" in i.rule_id]
        assert len(margin_issues) >= 1
        
        issue = margin_issues[0]
        assert issue.severity == Severity.ERROR
        assert "левое" in issue.message.lower() or "left" in issue.rule_id

    def test_all_margins_wrong_multiple_issues(self, tmp_path):
        """Проверка что все неправильные поля создают несколько Issue."""
        config = load_config()
        expected = config["margins_cm"]
        tol = expected.get("tolerance", 0.15)
        
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Cm(expected["left"] + tol + 0.5)
        section.right_margin = Cm(expected["right"] + tol + 0.5)
        section.top_margin = Cm(expected["top"] + tol + 0.5)
        section.bottom_margin = Cm(expected["bottom"] + tol + 0.5)
        
        doc.add_paragraph("Текст с неправильными полями")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = MarginsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Должно быть минимум 4 ошибки (по одному на каждое поле)
        assert result.error_count >= 4

    def test_multiple_sections_one_wrong_creates_issue(self, tmp_path):
        """Проверка что если одна секция имеет неправильные поля - создаётся Issue."""
        config = load_config()
        expected = config["margins_cm"]
        tol = expected.get("tolerance", 0.15)
        
        doc = Document()
        
        # Первая секция - корректная
        section1 = doc.sections[0]
        section1.left_margin = Cm(expected["left"])
        section1.right_margin = Cm(expected["right"])
        section1.top_margin = Cm(expected["top"])
        section1.bottom_margin = Cm(expected["bottom"])
        doc.add_paragraph("Секция 1")
        
        # Вторая секция - неправильная
        doc.add_section()
        section2 = doc.sections[1]
        section2.left_margin = Cm(expected["left"] + tol + 0.5)
        section2.right_margin = Cm(expected["right"])
        section2.top_margin = Cm(expected["top"])
        section2.bottom_margin = Cm(expected["bottom"])
        doc.add_paragraph("Секция 2")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = MarginsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Должна быть ошибка для второй секции
        assert result.error_count >= 1
        
        # В сообщении должна быть ссылка на секцию
        margin_issues = [i for i in result.issues if "margin" in i.rule_id]
        assert any("2" in (i.location_hint or "") for i in margin_issues)

    def test_tolerance_from_config_is_used(self, tmp_path):
        """Проверка что tolerance читается из конфига."""
        config = load_config()
        expected = config["margins_cm"]
        tol = expected.get("tolerance", 0.15)
        
        doc = Document()
        section = doc.sections[0]
        # Поле в пределах tolerance - не должно быть ошибки
        section.left_margin = Cm(expected["left"] + tol * 0.5)
        section.right_margin = Cm(expected["right"])
        section.top_margin = Cm(expected["top"])
        section.bottom_margin = Cm(expected["bottom"])
        
        doc.add_paragraph("Текст с полем в пределах tolerance")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = MarginsCheck(rules=config)
        result = check.run(model, resolver)
        
        # Не должно быть ошибок так как отклонение в пределах tolerance
        margin_issues = [i for i in result.issues if "margin_left" in i.rule_id]
        assert len(margin_issues) == 0


class TestSpacingCheckIntegration:
    """Интеграционные тесты проверки межстрочного интервала."""

    def test_correct_spacing_no_issues(self, tmp_path):
        """Проверка что корректный интервал не создаёт Issue."""
        config = load_config()
        
        doc = Document()
        doc.add_paragraph("Введение")
        para = doc.add_paragraph("Текст с правильным интервалом")
        
        # Устанавливаем интервал 1.5 через XML
        pPr = para._p.get_or_add_pPr()
        spacing = pPr.get_or_add_spacing()
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "360")
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule", "auto")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = SpacingCheck(rules=config)
        result = check.run(model, resolver)
        
        spacing_issues = [i for i in result.issues if i.rule_id == "line_spacing"]
        assert len(spacing_issues) == 0

    def test_wrong_spacing_creates_issue(self, tmp_path):
        """Проверка что неправильный интервал создаёт Issue."""
        config = load_config()
        
        doc = Document()
        doc.add_paragraph("Введение")
        para = doc.add_paragraph("Текст с одинарным интервалом")
        
        # Устанавливаем интервал 1.0 (нарушение, должно быть 1.5)
        pPr = para._p.get_or_add_pPr()
        spacing = pPr.get_or_add_spacing()
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "240")
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule", "auto")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = SpacingCheck(rules=config)
        result = check.run(model, resolver)
        
        spacing_issues = [i for i in result.issues if i.rule_id == "line_spacing"]
        # Интервал 1.0 вместо 1.5 должен создать ошибку
        assert len(spacing_issues) >= 1
        
        issue = spacing_issues[0]
        assert issue.severity == Severity.ERROR

    def test_spacing_before_intro_ignored(self, tmp_path):
        """Проверка что интервал до Введения не проверяется."""
        config = load_config()
        
        doc = Document()
        # Текст до Введения - не проверяется
        para1 = doc.add_paragraph("Текст до введения с неправильным интервалом")
        pPr = para1._p.get_or_add_pPr()
        spacing = pPr.get_or_add_spacing()
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "240")
        
        doc.add_paragraph("Введение")
        para2 = doc.add_paragraph("Текст после введения")
        pPr2 = para2._p.get_or_add_pPr()
        spacing2 = pPr2.get_or_add_spacing()
        spacing2.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "360")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = SpacingCheck(rules=config)
        result = check.run(model, resolver)
        
        # Ошибки должны быть только для текста после Введения
        spacing_issues = [i for i in result.issues if i.rule_id == "line_spacing"]
        assert len(spacing_issues) == 0


class TestIndentsCheckIntegration:
    """Интеграционные тесты проверки абзацных отступов."""

    def test_correct_body_indent_no_issues(self, tmp_path):
        """Проверка что корректный отступ не создаёт Issue."""
        config = load_config()
        body_indent = config["indents"]["body_first_line"]
        tol = config["indents"].get("tolerance", 0.1)
        
        doc = Document()
        doc.add_paragraph("Введение")
        para = doc.add_paragraph("Текст с правильным отступом")
        
        # Устанавливаем отступ первой строки 1.25 см (в twips: 1 см ≈ 567 twips)
        indent_twips = int(body_indent * 567)
        pPr = para._p.get_or_add_pPr()
        ind = pPr.get_or_add_ind()
        ind.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLine", str(indent_twips))
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = IndentsCheck(rules=config)
        result = check.run(model, resolver)
        
        indent_issues = [i for i in result.issues if "indent" in i.rule_id]
        assert len(indent_issues) == 0

    def test_heading_with_indent_creates_error(self, tmp_path):
        """Проверка что отступ у заголовка создаёт Error."""
        config = load_config()
        
        doc = Document()
        doc.add_paragraph("Введение")
        heading = doc.add_paragraph("Глава 1")
        
        # У заголовка НЕ должно быть отступа, но мы его добавим
        indent_twips = 700  # ~1.23 см
        pPr = heading._p.get_or_add_pPr()
        ind = pPr.get_or_add_ind()
        ind.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLine", str(indent_twips))
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = IndentsCheck(rules=config)
        result = check.run(model, resolver)
        
        heading_issues = [i for i in result.issues if i.rule_id == "heading_indent"]
        assert len(heading_issues) >= 1
        
        issue = heading_issues[0]
        assert issue.severity == Severity.ERROR

    def test_body_without_indent_creates_warning(self, tmp_path):
        """Проверка что отсутствие отступа у текста создаёт Warning."""
        config = load_config()
        
        doc = Document()
        doc.add_paragraph("Введение")
        para = doc.add_paragraph("Текст без отступа")
        # Отступ первой строки не задан (должен быть 1.25 см)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = IndentsCheck(rules=config)
        result = check.run(model, resolver)
        
        body_issues = [i for i in result.issues if i.rule_id == "body_indent"]
        # Отсутствие отступа должно создать warning
        assert len(body_issues) >= 1
        
        issue = body_issues[0]
        assert issue.severity == Severity.WARNING


class TestConfigLoading:
    """Тесты загрузки и использования конфигурации."""

    def test_config_loaded_successfully(self):
        """Проверка что конфиг загружается без ошибок."""
        config = load_config()
        
        assert config is not None
        assert "fonts" in config
        assert "margins_cm" in config
        assert "spacing" in config
        assert "indents" in config

    def test_config_has_required_font(self):
        """Проверка что в конфиге есть требуемый шрифт."""
        config = load_config()
        
        assert "required_font" in config["fonts"]
        assert config["fonts"]["required_font"] == "Times New Roman"

    def test_config_has_font_sizes(self):
        """Проверка что в конфиге есть размеры шрифтов."""
        config = load_config()
        
        assert config["fonts"]["sizes"]["body"] == 14
        assert 11 in config["fonts"]["sizes"]["table_cell"]
        assert 12 in config["fonts"]["sizes"]["table_cell"]

    def test_config_has_margins_with_tolerance(self):
        """Проверка что в конфиге есть поля с tolerance."""
        config = load_config()
        
        assert config["margins_cm"]["left"] == 3.0
        assert config["margins_cm"]["right"] == 1.0
        assert config["margins_cm"]["top"] == 2.0
        assert config["margins_cm"]["bottom"] == 2.0
        assert "tolerance" in config["margins_cm"]
        assert config["margins_cm"]["tolerance"] == 0.15

    def test_config_has_spacing_tolerance(self):
        """Проверка что в конфиге есть tolerance для интервалов."""
        config = load_config()
        
        assert config["spacing"]["body"] == 1.5
        assert config["spacing"]["tolerance"] == 0.05

    def test_config_has_indent_tolerance(self):
        """Проверка что в конфиге есть tolerance для отступов."""
        config = load_config()
        
        assert config["indents"]["body_first_line"] == 1.25
        assert config["indents"]["tolerance"] == 0.1


class TestCheckResultStructure:
    """Тесты структуры CheckResult."""

    def test_check_result_has_check_id(self, tmp_path):
        """Проверка что CheckResult содержит check_id."""
        config = load_config()
        
        doc = Document()
        doc.add_paragraph("Тест")
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        assert result.check_id == "fonts"
        assert result.check_name == "Шрифт и кегль"

    def test_check_result_counts_errors(self, tmp_path):
        """Проверка что error_count подсчитывает правильно."""
        config = load_config()
        
        doc = Document()
        # Создаём несколько параграфов с нарушением
        for i in range(10):
            para = doc.add_paragraph(f"Текст {i} шрифтом Arial")
            para.runs[0].font.name = "Arial"
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        # error_count должен совпадать с количеством Issue severity ERROR
        error_count = sum(1 for i in result.issues if i.severity == Severity.ERROR)
        assert result.error_count == error_count

    def test_issue_has_required_fields(self, tmp_path):
        """Проверка что Issue содержит все обязательные поля."""
        config = load_config()
        
        doc = Document()
        # Добавляем Введение чтобы проверки работали
        doc.add_paragraph("Введение")
        para = doc.add_paragraph("Текст шрифтом Arial")
        para.runs[0].font.name = "Arial"
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        resolver = StyleResolver(model.doc)
        
        check = FontsCheck(rules=config)
        result = check.run(model, resolver)
        
        assert len(result.issues) > 0
        issue = result.issues[0]
        
        assert issue.rule_id
        assert issue.message
        assert isinstance(issue.severity, Severity)
        # location_hint и context могут быть пустыми
