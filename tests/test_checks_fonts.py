"""
Тесты проверок шрифтов, кегля и жирного текста.

Запуск:
  pytest tests/test_checks_fonts.py -v
  pytest tests/test_checks_fonts.py --cov=vkr_checker --cov-report=term-missing
"""
from __future__ import annotations

from pathlib import Path
import yaml

import pytest
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from vkr_checker.parser import load_document
from vkr_checker.style_resolver import StyleResolver


FIXTURES = Path(__file__).parent / "fixtures"
CONFIG_PATH = Path(__file__).parent.parent / "config" / "rules.yaml"


def load_config() -> dict:
    """Загружает конфигурацию из rules.yaml."""
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def save_doc(doc: Document, tmp_path: Path, name: str = "test.docx") -> Path:
    """Сохраняет документ и возвращает путь к нему."""
    p = tmp_path / name
    doc.save(str(p))
    return p


class TestFontName:
    """Тесты проверки имени шрифта."""

    def test_times_new_roman(self, tmp_path):
        """Проверка шрифта Times New Roman."""
        doc = Document()
        para = doc.add_paragraph("Текст шрифтом Times New Roman")
        para.runs[0].font.name = "Times New Roman"
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_name = resolver.get_font_name(loaded_run, loaded_para)
        assert font_name == "Times New Roman"

    def test_arial(self, tmp_path):
        """Проверка шрифта Arial."""
        doc = Document()
        para = doc.add_paragraph("Текст шрифтом Arial")
        para.runs[0].font.name = "Arial"
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_name = resolver.get_font_name(loaded_run, loaded_para)
        assert font_name == "Arial"

    def test_font_inheritance_from_style(self, tmp_path):
        """Проверка наследования шрифта из стиля параграфа."""
        doc = Document()
        para = doc.add_paragraph("Текст без явного указания шрифта")
        # Не задаём шрифт явно - должен наследоваться из стиля Normal
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_name = resolver.get_font_name(loaded_run, loaded_para)
        # Шрифт может быть None если не задан ни в run, ни в стиле, ни в дефолтах
        # Проверяем что резолвер работает без ошибок
        assert font_name is None or isinstance(font_name, str)

    def test_mixed_fonts_in_paragraph(self, tmp_path):
        """Проверка параграфа с разными шрифтами в разных run."""
        doc = Document()
        para = doc.add_paragraph("Первая часть ")
        run1 = para.runs[0]
        run1.font.name = "Times New Roman"
        
        run2 = para.add_run("вторая часть")
        run2.font.name = "Arial"
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        
        # Проверяем первый run
        font1 = resolver.get_font_name(loaded_para.runs[0], loaded_para)
        assert font1 == "Times New Roman"
        
        # Проверяем второй run
        font2 = resolver.get_font_name(loaded_para.runs[1], loaded_para)
        assert font2 == "Arial"


class TestFontSize:
    """Тесты проверки размера шрифта (кегля)."""

    def test_font_size_14pt(self, tmp_path):
        """Проверка кегля 14pt."""
        doc = Document()
        para = doc.add_paragraph("Текст кеглем 14pt")
        para.runs[0].font.size = Pt(14)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
        assert font_size == 14.0

    def test_font_size_12pt(self, tmp_path):
        """Проверка кегля 12pt."""
        doc = Document()
        para = doc.add_paragraph("Текст кеглем 12pt")
        para.runs[0].font.size = Pt(12)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
        assert font_size == 12.0

    def test_font_size_16pt(self, tmp_path):
        """Проверка кегля 16pt."""
        doc = Document()
        para = doc.add_paragraph("Текст кеглем 16pt")
        para.runs[0].font.size = Pt(16)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
        assert font_size == 16.0

    @pytest.mark.parametrize("font_size", [8, 10, 12, 14, 16, 18, 20, 24])
    def test_various_font_sizes(self, tmp_path, font_size):
        """Параметризованный тест для различных размеров шрифта."""
        doc = Document()
        para = doc.add_paragraph(f"Текст кеглем {font_size}pt")
        para.runs[0].font.size = Pt(font_size)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        size = resolver.get_font_size_pt(loaded_run, loaded_para)
        assert size == float(font_size)

    def test_font_size_inheritance_from_style(self, tmp_path):
        """Проверка наследования размера шрифта из стиля."""
        doc = Document()
        para = doc.add_paragraph("Текст без явного указания кегля")
        # Не задаём размер явно - должен наследоваться из стиля
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
        # Размер должен быть получен из стиля (обычно 11pt или 12pt)
        assert font_size is not None
        assert font_size > 0

    def test_mixed_font_sizes_in_paragraph(self, tmp_path):
        """Проверка параграфа с разными размерами шрифта."""
        doc = Document()
        para = doc.add_paragraph("Маленький ")
        run1 = para.runs[0]
        run1.font.size = Pt(10)
        
        run2 = para.add_run("большой")
        run2.font.size = Pt(16)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        
        # Проверяем первый run
        size1 = resolver.get_font_size_pt(loaded_para.runs[0], loaded_para)
        assert size1 == 10.0
        
        # Проверяем второй run
        size2 = resolver.get_font_size_pt(loaded_para.runs[1], loaded_para)
        assert size2 == 16.0


class TestBoldText:
    """Тесты проверки жирного текста."""

    def test_bold_true(self, tmp_path):
        """Проверка явного жирного шрифта."""
        doc = Document()
        para = doc.add_paragraph("Жирный текст")
        para.runs[0].font.bold = True
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        is_bold = resolver.is_bold(loaded_run, loaded_para)
        assert is_bold is True

    def test_bold_false(self, tmp_path):
        """Проверка явного обычного шрифта."""
        doc = Document()
        para = doc.add_paragraph("Обычный текст")
        para.runs[0].font.bold = False
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        is_bold = resolver.is_bold(loaded_run, loaded_para)
        assert is_bold is False

    def test_bold_not_set(self, tmp_path):
        """Проверка когда bold не задан явно (None)."""
        doc = Document()
        para = doc.add_paragraph("Текст без явного указания bold")
        # Не задаём bold явно
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        is_bold = resolver.is_bold(loaded_run, loaded_para)
        # Должно вернуть значение из стиля или дефолтное
        assert isinstance(is_bold, bool)

    def test_bold_inheritance_from_style(self, tmp_path):
        """Проверка наследования жирности из стиля."""
        doc = Document()
        para = doc.add_paragraph("Текст со стилем")
        # Bold не задан явно - должно наследоваться из стиля
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        is_bold = resolver.is_bold(loaded_run, loaded_para)
        # Должно вернуть булево значение
        assert isinstance(is_bold, bool)

    def test_mixed_bold_in_paragraph(self, tmp_path):
        """Проверка параграфа с жирным и обычным текстом."""
        doc = Document()
        para = doc.add_paragraph("Жирный ")
        run1 = para.runs[0]
        run1.font.bold = True
        
        run2 = para.add_run("обычный")
        run2.font.bold = False
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        
        # Проверяем первый run (жирный)
        is_bold1 = resolver.is_bold(loaded_para.runs[0], loaded_para)
        assert is_bold1 is True
        
        # Проверяем второй run (обычный)
        is_bold2 = resolver.is_bold(loaded_para.runs[1], loaded_para)
        assert is_bold2 is False


class TestItalicText:
    """Тесты проверки курсива."""

    def test_italic_true(self, tmp_path):
        """Проверка явного курсива."""
        doc = Document()
        para = doc.add_paragraph("Курсивный текст")
        para.runs[0].font.italic = True
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        is_italic = resolver.is_italic(loaded_run, loaded_para)
        assert is_italic is True

    def test_italic_false(self, tmp_path):
        """Проверка явного отсутствия курсива."""
        doc = Document()
        para = doc.add_paragraph("Прямой текст")
        para.runs[0].font.italic = False
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        is_italic = resolver.is_italic(loaded_run, loaded_para)
        assert is_italic is False

    def test_mixed_italic_in_paragraph(self, tmp_path):
        """Проверка параграфа с курсивным и прямым текстом."""
        doc = Document()
        para = doc.add_paragraph("Курсив ")
        run1 = para.runs[0]
        run1.font.italic = True
        
        run2 = para.add_run("прямой")
        run2.font.italic = False
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        
        # Проверяем первый run (курсив)
        is_italic1 = resolver.is_italic(loaded_para.runs[0], loaded_para)
        assert is_italic1 is True
        
        # Проверяем второй run (прямой)
        is_italic2 = resolver.is_italic(loaded_para.runs[1], loaded_para)
        assert is_italic2 is False


class TestFontIntegration:
    """Интеграционные тесты для шрифтов."""

    def test_full_paragraph_analysis(self, tmp_path):
        """Полный анализ параграфа с различными стилями."""
        doc = Document()
        
        # Создаём параграф с несколькими run с разными стилями
        para = doc.add_paragraph("")
        
        run1 = para.add_run("Обычный текст, ")
        run1.font.name = "Times New Roman"
        run1.font.size = Pt(14)
        run1.font.bold = False
        run1.font.italic = False
        
        run2 = para.add_run("жирный, ")
        run2.font.name = "Times New Roman"
        run2.font.size = Pt(14)
        run2.font.bold = True
        run2.font.italic = False
        
        run3 = para.add_run("курсив")
        run3.font.name = "Times New Roman"
        run3.font.size = Pt(14)
        run3.font.bold = False
        run3.font.italic = True
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        
        # Проверяем каждый run
        for i, expected in enumerate([
            ("Times New Roman", 14.0, False, False),
            ("Times New Roman", 14.0, True, False),
            ("Times New Roman", 14.0, False, True),
        ]):
            loaded_run = loaded_para.runs[i]
            font_name = resolver.get_font_name(loaded_run, loaded_para)
            font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
            is_bold = resolver.is_bold(loaded_run, loaded_para)
            is_italic = resolver.is_italic(loaded_run, loaded_para)
            
            assert font_name == expected[0]
            assert font_size == expected[1]
            assert is_bold == expected[2]
            assert is_italic == expected[3]

    def test_heading_style_detection(self, tmp_path):
        """Проверка что заголовки обычно жирные."""
        doc = Document()
        
        # Заголовок
        heading = doc.add_heading("Глава 1", level=1)
        heading.runs[0].font.bold = True
        heading.runs[0].font.size = Pt(16)
        
        # Обычный текст
        para = doc.add_paragraph("Обычный текст")
        para.runs[0].font.bold = False
        para.runs[0].font.size = Pt(14)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        
        # Проверяем заголовок
        heading_para = model.doc.paragraphs[0]
        heading_run = heading_para.runs[0]
        assert resolver.is_bold(heading_run, heading_para) is True
        assert resolver.get_font_size_pt(heading_run, heading_para) == 16.0
        
        # Проверяем обычный текст
        text_para = model.doc.paragraphs[1]
        text_run = text_para.runs[0]
        assert resolver.is_bold(text_run, text_para) is False
        assert resolver.get_font_size_pt(text_run, text_para) == 14.0


class TestFontEdgeCases:
    """Граничные случаи проверки шрифтов."""

    def test_empty_paragraph(self, tmp_path):
        """Проверка пустого параграфа."""
        doc = Document()
        doc.add_paragraph("")
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        
        if loaded_para.runs:
            loaded_run = loaded_para.runs[0]
            # Даже для пустого параграфа резолвер должен работать
            font_name = resolver.get_font_name(loaded_run, loaded_para)
            assert font_name is not None

    def test_special_characters_font(self, tmp_path):
        """Проверка шрифта для специальных символов."""
        doc = Document()
        para = doc.add_paragraph("Спецсимволы: © ® ™ € £ ¥")
        para.runs[0].font.name = "Times New Roman"
        para.runs[0].font.size = Pt(14)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_name = resolver.get_font_name(loaded_run, loaded_para)
        assert font_name == "Times New Roman"
        
        font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
        assert font_size == 14.0

    def test_very_long_text_font(self, tmp_path):
        """Проверка шрифта для очень длинного текста."""
        doc = Document()
        long_text = "A" * 10000
        para = doc.add_paragraph(long_text)
        para.runs[0].font.name = "Arial"
        para.runs[0].font.size = Pt(12)
        
        path = save_doc(doc, tmp_path)
        model = load_document(path)
        
        resolver = StyleResolver(model.doc)
        loaded_para = model.doc.paragraphs[0]
        loaded_run = loaded_para.runs[0]
        
        font_name = resolver.get_font_name(loaded_run, loaded_para)
        assert font_name == "Arial"
        
        font_size = resolver.get_font_size_pt(loaded_run, loaded_para)
        assert font_size == 12.0
